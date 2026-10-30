"""Parse a .docx report into an ordered list of content blocks.

Block shapes (all have a "type"):
  heading   {level: 1-3, text}
  paragraph {html, text, style: "lede"|"note"|"caption"|None}
  list      {ordered: bool, items: [html, ...]}
  image     {src, caption, exhibit, width, height, note}
  table     {rows: [[cell, ...], ...], header: bool, caption, exhibit, note}
  layout    {cols: n, cells: [[block, ...], ...]}   side-by-side figures
  callout   {text}
  pagebreak {}

Real AID reports mostly use the Normal style with manual formatting, so
headings are also detected from formatting (bold, short, larger font), and
the masthead lines at the top (brand, title, coverage) are lifted into meta.
"""
from __future__ import annotations

import html
import re
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

from .common import attach_captions, build_exhibits, guess_language, is_caption, keywords_from_coverage

IMAGE_EXT = {
    "image/png": "png", "image/jpeg": "jpg", "image/jpg": "jpg", "image/gif": "gif", "image/bmp": "bmp",
    "image/webp": "webp", "image/svg+xml": "svg", "image/tiff": "tiff", "image/x-emf": "emf", "image/x-wmf": "wmf",
}
HEADING_RE = re.compile(r"^(?:heading|标题|título|titre)\s*(\d)", re.I)
LIST_RE = re.compile(r"list|列表|项目符号|编号", re.I)
TITLE_RE = re.compile(r"^(?:title|标题$|subtitle|副标题)", re.I)
NOTE_RE = re.compile(r"^(source|sources|资料来源|来源|note|notes|注|注释)\s*[:：]", re.I)
BULLET_RE = re.compile(r"^[•·▪◦●○■□➢➤►▶\-–—]\s+")
REPORT_KW = re.compile(r"report|update|summary|quarterly|tracker|review|outlook|报告|更新|摘要|季度|追踪|展望", re.I)
COVERAGE_RE = re.compile(r"^(coverage|覆盖范围)\s*[:：]\s*", re.I)
END_PUNCT = ("。", ".", "!", "?", "！", "？", ";", "；", ":", "：", ",", "，")


def _convert_to_png(blob: bytes) -> bytes | None:
    try:
        from PIL import Image  # type: ignore

        im = Image.open(BytesIO(blob))
        out = BytesIO()
        im.convert("RGBA").save(out, format="PNG")
        return out.getvalue()
    except Exception:
        return None


class DocxParser:
    def __init__(self, path: Path, image_dir: Path):
        self.path = Path(path)
        self.image_dir = Path(image_dir)
        self.image_dir.mkdir(parents=True, exist_ok=True)
        self.doc = Document(str(self.path))
        self.blocks: list[dict[str, Any]] = []
        self._img_n = 0
        self._list_buf: list[str] | None = None
        self._list_ordered = False
        self._fmt_levels: dict[float, int] = {}
        self._style_levels_used = 0
        self.meta: dict[str, Any] = {"masthead": "", "title": "", "subtitle": "", "coverage": "", "publish_date": ""}
        self._skip: set = set()  # paragraph elements lifted into meta
        self._in_masthead = True

    # ------------------------------------------------------------------ public
    def parse(self) -> dict[str, Any]:
        self._prescan()
        self._walk(self.doc.element.body, self.blocks)
        self._flush_list(self.blocks)

        blocks = attach_captions(self.blocks)
        for b in blocks:
            if b["type"] == "layout":
                b["cells"] = [attach_captions(c) for c in b["cells"]]
        exhibits = build_exhibits(blocks)
        text = " ".join(_block_text(b) for b in blocks)
        props = self.doc.core_properties
        title = (props.title or "").strip() or self.meta["title"] or self._first_heading(blocks) or self.path.stem
        keywords = [k.strip() for k in re.split(r"[;,，；]", props.keywords or "") if k.strip()]
        if not keywords and self.meta["coverage"]:
            keywords = keywords_from_coverage(self.meta["coverage"])
        return {
            "title": title,
            "subtitle": self.meta["subtitle"],
            "masthead": self.meta["masthead"],
            "coverage": self.meta["coverage"],
            "publish_date": self.meta["publish_date"],
            "keywords": keywords,
            "language": guess_language(text),
            "blocks": blocks,
            "exhibits": exhibits,
            "stats": {
                "blocks": len(blocks),
                "images": len([e for e in exhibits if e["type"] == "image"]),
                "tables": len([e for e in exhibits if e["type"] == "table"]),
                "chars": len(text),
            },
        }

    # ----------------------------------------------------------------- prescan
    def _prescan(self):
        """Lift masthead lines into meta and learn which font sizes act as headings."""
        body = self.doc.element.body
        paras = [Paragraph(p, self.doc) for p in body.iterchildren() if p.tag == qn("w:p")]
        leading = [p for p in paras if p.text.strip()][:6]
        for p in leading:
            txt = p.text.strip()
            size = self._max_size(p)
            if self._style_level(p):
                break
            if re.search(r"junson", txt, re.I) and not REPORT_KW.search(txt.replace("Research", "").replace("研究", "")):
                self.meta["masthead"] = txt.split("\t")[0].strip()
                self._skip.add(p._p)
                continue
            if COVERAGE_RE.match(txt):
                parts = [x.strip() for x in COVERAGE_RE.sub("", txt).split("\t") if x.strip()]
                self.meta["coverage"] = parts[0] if parts else ""
                if len(parts) > 1:
                    self.meta["publish_date"] = parts[1]
                self._skip.add(p._p)
                continue
            if not self.meta["title"] and (REPORT_KW.search(txt) or (size and size >= 14)) and len(txt) <= 160:
                self.meta["title"] = re.sub(r"\s{2,}", " ", txt.split("\t")[0]).strip()
                self._skip.add(p._p)
                continue
            if self.meta["title"] and REPORT_KW.search(txt) and not REPORT_KW.search(self.meta["title"]) and len(txt) <= 160:
                # a big company/brand line came first; the real title names the report
                if not self.meta["masthead"]:
                    self.meta["masthead"] = self.meta["title"]
                self.meta["title"] = re.sub(r"\s{2,}", " ", txt.split("\t")[0]).strip()
                self._skip.add(p._p)
                continue
            if not self.meta["masthead"] and size and size >= 20 and len(txt) <= 120:
                self.meta["masthead"] = txt.split("\t")[0].strip()
                self._skip.add(p._p)
                continue
            if self.meta["title"] and not self.meta["subtitle"] and REPORT_KW.search(txt) and len(txt) <= 160:
                self.meta["subtitle"] = txt.strip()
                self._skip.add(p._p)
                continue
            break

        sizes: set[float] = set()
        style_levels: set[int] = set()
        for p in paras:
            if p._p in self._skip or not p.text.strip():
                continue
            lvl = self._style_level(p)
            if lvl:
                style_levels.add(lvl)
                continue
            s = self._fmt_heading_size(p)
            if s:
                sizes.add(s)
        self._style_levels_used = max(style_levels) if style_levels else 0
        ordered = sorted(sizes, reverse=True)
        base = self._style_levels_used
        for i, s in enumerate(ordered):
            self._fmt_levels[s] = min(3, base + i + 1)

    # ----------------------------------------------------------------- helpers
    @staticmethod
    def _first_heading(blocks: list[dict[str, Any]]) -> str:
        for b in blocks:
            if b["type"] == "heading":
                return b["text"]
        return ""

    def _style_name(self, p: Paragraph) -> str:
        try:
            return (p.style.name or "") if p.style is not None else ""
        except Exception:
            return ""

    def _style_level(self, p: Paragraph) -> int:
        """Heading level from the paragraph style (or outline level), 0 if none."""
        name = self._style_name(p)
        m = HEADING_RE.match(name)
        if m:
            return max(1, min(3, int(m.group(1))))
        st = p.style
        while st is not None:
            m = HEADING_RE.match(st.name or "")
            if m:
                return max(1, min(3, int(m.group(1))))
            st = st.base_style
        pPr = p._p.pPr
        if pPr is not None:
            ol = pPr.find(qn("w:outlineLvl"))
            if ol is not None:
                try:
                    return max(1, min(3, int(ol.get(qn("w:val"))) + 1))
                except Exception:
                    pass
        return 0

    @staticmethod
    def _max_size(p: Paragraph) -> float | None:
        sizes = [r.font.size.pt for r in p.runs if r.text.strip() and r.font.size]
        return max(sizes) if sizes else None

    def _fmt_heading_size(self, p: Paragraph) -> float | None:
        """Return a pseudo size if the paragraph looks like a manually formatted heading."""
        txt = p.text.strip()
        if not txt or len(txt) > 110 or "\n" in txt:
            return None
        if is_caption(txt) or NOTE_RE.match(txt) or BULLET_RE.match(txt):
            return None
        if txt.endswith(END_PUNCT):
            return None
        if self._is_list(p)[0]:
            return None
        runs = [r for r in p.runs if r.text.strip()]
        if not runs:
            return None
        bold = all(r.bold or (r.style is not None and r.style.font.bold) for r in runs)
        if not bold:
            return None
        size = self._max_size(p)
        if size is None:
            return 0.0  # bold, short, no explicit size: lowest heading level
        if size < 10.5:
            return None
        return size

    def _is_list(self, p: Paragraph) -> tuple[bool, bool]:
        pPr = p._p.pPr
        if pPr is not None and pPr.find(qn("w:numPr")) is not None:
            return True, self._num_is_ordered(pPr.find(qn("w:numPr")))
        name = self._style_name(p)
        if LIST_RE.search(name) and name.lower() != "list paragraph":
            return True, bool(re.search(r"number|编号", name, re.I))
        return False, False

    def _num_is_ordered(self, num_pr) -> bool:
        try:
            num_id_el = num_pr.find(qn("w:numId"))
            ilvl_el = num_pr.find(qn("w:ilvl"))
            if num_id_el is None:
                return False
            num_id = num_id_el.get(qn("w:val"))
            ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
            numbering = self.doc.part.numbering_part.element
            for num in numbering.findall(qn("w:num")):
                if num.get(qn("w:numId")) == num_id:
                    abs_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
                    for abs_num in numbering.findall(qn("w:abstractNum")):
                        if abs_num.get(qn("w:abstractNumId")) == abs_id:
                            for lvl in abs_num.findall(qn("w:lvl")):
                                if lvl.get(qn("w:ilvl")) == ilvl:
                                    fmt = lvl.find(qn("w:numFmt"))
                                    val = fmt.get(qn("w:val")) if fmt is not None else "bullet"
                                    return val not in ("bullet", "none")
        except Exception:
            pass
        return False

    def _runs_html(self, p: Paragraph) -> tuple[str, str]:
        parts: list[str] = []
        texts: list[str] = []
        for child in p._p.iterchildren():
            if child.tag == qn("w:r"):
                self._run_html(child, parts, texts)
            elif child.tag == qn("w:hyperlink"):
                for r in child.findall(qn("w:r")):
                    self._run_html(r, parts, texts)
            elif child.tag in (qn("w:ins"), qn("w:smartTag"), qn("w:sdt")):
                for r in child.iter(qn("w:r")):
                    self._run_html(r, parts, texts)
        return "".join(parts), "".join(texts)

    def _run_html(self, r_el, parts: list[str], texts: list[str]):
        t = "".join((x.text or "") if x.tag == qn("w:t") else ("\n" if x.tag == qn("w:br") else ("\t" if x.tag == qn("w:tab") else "")) for x in r_el.iterchildren())
        if not t:
            return
        texts.append(t)
        seg = html.escape(t).replace("\n", "<br>").replace("\t", " ")
        rPr = r_el.find(qn("w:rPr"))
        if rPr is not None:
            def on(tag):
                e = rPr.find(qn(tag))
                return e is not None and e.get(qn("w:val")) not in ("0", "false", "none")
            if on("w:b"):
                seg = f"<strong>{seg}</strong>"
            if on("w:i"):
                seg = f"<em>{seg}</em>"
            if on("w:u"):
                seg = f"<u>{seg}</u>"
            va = rPr.find(qn("w:vertAlign"))
            if va is not None:
                v = va.get(qn("w:val"))
                if v == "superscript":
                    seg = f"<sup>{seg}</sup>"
                elif v == "subscript":
                    seg = f"<sub>{seg}</sub>"
        parts.append(seg)

    def _extract_images(self, p: Paragraph) -> list[dict[str, Any]]:
        out: list[dict[str, Any]] = []
        for blip in p._p.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed")) or blip.get(qn("r:link"))
            if not rid:
                continue
            try:
                part = self.doc.part.related_parts[rid]
            except KeyError:
                continue
            blob = part.blob
            ctype = getattr(part, "content_type", "") or ""
            ext = IMAGE_EXT.get(ctype, "bin")
            if ext in ("bmp", "tiff", "emf", "wmf", "bin"):
                conv = _convert_to_png(blob)
                if conv:
                    blob, ext = conv, "png"
            width = height = None
            ext_el = next(iter(p._p.iter(qn("wp:extent"))), None)
            if ext_el is not None:
                try:
                    width = int(ext_el.get("cx")) / 9525
                    height = int(ext_el.get("cy")) / 9525
                except Exception:
                    pass
            self._img_n += 1
            name = f"img_{self._img_n:03d}.{ext}"
            (self.image_dir / name).write_bytes(blob)
            out.append({"type": "image", "src": name, "caption": "", "exhibit": None, "width": width, "height": height, "note": "", "unsupported": ext in ("emf", "wmf", "bin")})
        return out

    # ------------------------------------------------------------------- walk
    def _walk(self, container, target: list[dict[str, Any]]):
        for child in container.iterchildren():
            tag = child.tag
            if tag == qn("w:p"):
                self._handle_paragraph(Paragraph(child, self.doc), target)
            elif tag == qn("w:tbl"):
                self._flush_list(target)
                self._handle_table(Table(child, self.doc), target)
            elif tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    self._walk(content, target)

    def _flush_list(self, target: list[dict[str, Any]]):
        if self._list_buf:
            target.append({"type": "list", "ordered": self._list_ordered, "items": self._list_buf})
        self._list_buf = None

    def _push(self, target, block):
        self._flush_list(target)
        target.append(block)

    def _handle_paragraph(self, p: Paragraph, target: list[dict[str, Any]]):
        if p._p in self._skip:
            return
        if p._p.find(".//" + qn("w:br") + "[@" + qn("w:type") + "='page']") is not None:
            self._push(target, {"type": "pagebreak"})

        images = self._extract_images(p)
        inner_html, text = self._runs_html(p)
        text = text.strip()
        style = self._style_name(p)

        if images:
            if self._in_masthead and not text:
                # logo above the title: not an exhibit
                for im in images:
                    try:
                        (self.image_dir / im["src"]).unlink()
                    except Exception:
                        pass
                    self._img_n -= 1
                return
            self._in_masthead = False
            self._flush_list(target)
            for im in images:
                target.append(im)
            if text:
                target.append({"type": "paragraph", "html": inner_html, "text": text, "style": "caption" if is_caption(text) else None})
            return

        if not text:
            self._flush_list(target)
            return
        self._in_masthead = False

        level = self._style_level(p)
        if level:
            self._push(target, {"type": "heading", "level": level, "text": text})
            return
        if TITLE_RE.match(style):
            self._push(target, {"type": "heading", "level": 1, "text": text, "doc_title": True})
            return

        is_list, ordered = self._is_list(p)
        literal = BULLET_RE.match(text)
        if literal:
            is_list, ordered = True, False
            inner_html = BULLET_RE.sub("", inner_html, count=1) if inner_html.startswith(literal.group(0).strip()[:1]) else inner_html
            # bullet may sit in its own run; strip it from the html too
            inner_html = re.sub(r"^(<[^>]+>)*[•·▪◦●○■□➢➤►▶\-–—]\s*", lambda m: (m.group(1) or ""), inner_html, count=1)
        elif re.match(r"^\(?\d{1,2}[.)]\s+\S", text) and len(text) < 400 and not self._fmt_heading_size(p):
            pass  # numbered sentences stay paragraphs; Word lists are detected via numPr
        if is_list:
            if self._list_buf is None or self._list_ordered != ordered:
                self._flush_list(target)
                self._list_buf, self._list_ordered = [], ordered
            self._list_buf.append(inner_html)
            return

        fs = self._fmt_heading_size(p)
        if fs is not None:
            lvl = self._fmt_levels.get(fs, min(3, self._style_levels_used + 1))
            self._push(target, {"type": "heading", "level": lvl, "text": text})
            return

        self._flush_list(target)
        pstyle = None
        if re.search(r"caption|题注", style, re.I) or is_caption(text):
            pstyle = "caption"
        elif NOTE_RE.match(text) or re.search(r"^note|注释", style, re.I) or (self._max_size(p) or 99) <= 7.5:
            pstyle = "note"
        target.append({"type": "paragraph", "html": inner_html, "text": text, "style": pstyle})

    def _handle_table(self, t: Table, target: list[dict[str, Any]]):
        self._in_masthead = False
        has_img = t._tbl.find(".//" + qn("a:blip")) is not None
        nrows, ncols = len(t.rows), len(t.columns)
        if has_img or (nrows == 1 and ncols <= 3 and ncols > 1):
            # side-by-side figures: parse every cell as its own block list
            cells: list[list[dict[str, Any]]] = []
            for row in t.rows:
                last_tc = None
                for cell in row.cells:
                    if cell._tc is last_tc:
                        continue
                    last_tc = cell._tc
                    sub: list[dict[str, Any]] = []
                    saved = (self._list_buf, self._list_ordered)
                    self._list_buf = None
                    self._walk(cell._tc, sub)
                    self._flush_list(sub)
                    self._list_buf, self._list_ordered = saved
                    if sub:
                        cells.append(sub)
            if cells:
                target.append({"type": "layout", "cols": max(1, min(3, len(cells) if nrows == 1 else ncols)), "cells": cells})
            return
        if nrows == 1 and ncols == 1:
            txt = "\n".join(p.text.strip() for p in t.rows[0].cells[0].paragraphs if p.text.strip())
            if txt:
                target.append({"type": "callout", "text": txt})
            return
        rows: list[list[str]] = []
        for row in t.rows:
            cells_txt: list[str] = []
            last_tc = None
            for cell in row.cells:
                if cell._tc is last_tc:
                    continue
                last_tc = cell._tc
                cells_txt.append("\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip()))
            if any(c for c in cells_txt):
                rows.append(cells_txt)
        if not rows:
            return
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        target.append({"type": "table", "rows": rows, "header": True, "caption": "", "exhibit": None, "note": ""})


def _block_text(b: dict[str, Any]) -> str:
    if b["type"] in ("paragraph", "heading", "callout"):
        return b.get("text", "")
    if b["type"] == "list":
        return " ".join(re.sub(r"<[^>]+>", "", i) for i in b.get("items", []))
    if b["type"] == "layout":
        return " ".join(_block_text(x) for c in b["cells"] for x in c)
    return ""


def parse_docx(path: Path, image_dir: Path) -> dict[str, Any]:
    return DocxParser(path, image_dir).parse()
